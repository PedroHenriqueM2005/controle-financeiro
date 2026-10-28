"""Testes dos gráficos SVG e das exportações (CSV/XLSX) — incluindo round-trip."""

from __future__ import annotations

import io

from app.services import charts, compare, exporter, importer, ledger
from app.services.importer import ImportOptions

SERIE = [
    {"period": "2026-08", "label": "08/26", "entradas": 300000, "saidas": 250000, "saldo": 50000, "vazio": False},
    {"period": "2026-09", "label": "09/26", "entradas": 0, "saidas": 0, "saldo": 0, "vazio": True},
    {"period": "2026-10", "label": "10/26", "entradas": 350000, "saidas": 410000, "saldo": -60000, "vazio": False},
]

CATEGORIAS = [
    {"name": "Moradia", "total": 120000, "color": "#f97316"},
    {"name": "Mercado", "total": 60000, "color": "#ef4444"},
    {"name": "Zerada", "total": 0, "color": "#000"},
]


class TestCharts:
    def test_barras(self):
        svg = charts.bars_series(SERIE, title="Teste")
        assert svg.startswith("<svg") and svg.endswith("</svg>")
        assert svg.count("<rect") >= 6          # 2 barras por mês
        assert "Teste" in svg
        assert "10/26" in svg

    def test_barras_vazias(self):
        assert "Sem dados" in charts.bars_series([])

    def test_escape_de_html(self):
        svg = charts.bars_series([{**SERIE[0], "label": "<script>x</script>"}])
        assert "<script>" not in svg
        assert "&lt;script&gt;" in svg

    def test_hbar_ignora_zeros(self):
        svg = charts.hbar(CATEGORIAS, title="Categorias")
        assert "Moradia" in svg and "Zerada" not in svg

    def test_hbar_sem_dados(self):
        assert "Sem dados" in charts.hbar([{"name": "a", "total": 0}])

    def test_donut(self):
        svg = charts.donut(CATEGORIAS, center_label="gasto", center_value="R$ 1.800,00")
        assert svg.count("<circle") == 3        # trilha + 2 fatias
        assert "stroke-dasharray" in svg
        assert "R$ 1.800,00" in svg

    def test_donut_vazio(self):
        assert "Sem dados" in charts.donut([])

    def test_linha(self):
        pontos = [{"label": "08/26", "acumulado": 50000}, {"label": "09/26", "acumulado": -20000},
                  {"label": "10/26", "acumulado": 130000}]
        svg = charts.line_chart(pontos, title="Saldo acumulado")
        assert "<polyline" in svg and "<polygon" in svg
        assert svg.count("<circle") == 3

    def test_compact_brl(self):
        assert charts.compact_brl(123456789) == "R$ 1,2 mi"
        assert charts.compact_brl(1234500) == "R$ 12,3 mil"
        assert charts.compact_brl(99900) == "R$ 999"
        assert charts.compact_brl(-1500000) == "-R$ 15,0 mil"


class TestExportCSV:
    def _popula(self, conn):
        ledger.create_transaction(conn, {"type": "entrada", "amount": "3.500,00", "occurred_on": "2026-10-05",
                                         "description": "Salário", "category_name": "Salário", "method": "Pix"})
        ledger.create_transaction(conn, {"type": "saida", "amount": "1.200,00", "occurred_on": "2026-10-06",
                                         "description": "Aluguel", "category_name": "Moradia",
                                         "item_name": "Aluguel", "recurring": True})
        ledger.create_transaction(conn, {"type": "saida", "status": "previsto", "amount": "99,90",
                                         "occurred_on": "2026-10-25", "description": "Internet",
                                         "category_name": "Contas (água, luz, internet)"})

    def test_csv_do_mes(self, conn):
        self._popula(conn)
        nome, conteudo = exporter.csv_bytes(conn, "2026-10")
        assert nome.startswith("lancamentos-2026-10")
        texto = conteudo.decode("utf-8-sig")
        linhas = texto.splitlines()
        assert linhas[0].split(";")[0] == "Data"
        assert len(linhas) == 4
        assert "3500,00" in texto and "1200,00" in texto
        assert "Previsto" in texto and "Sim" in texto

    def test_csv_completo(self, conn):
        self._popula(conn)
        nome, _ = exporter.csv_bytes(conn, None)
        assert "completo" in nome

    def test_resumo_csv(self, conn):
        self._popula(conn)
        nome, conteudo = exporter.resumo_csv_bytes(conn, months=3)
        texto = conteudo.decode("utf-8-sig")
        assert nome.startswith("resumo-mensal")
        assert "Moradia" in texto
        assert "Entradas realizadas" in texto and "Saldo" in texto


class TestExportXLSX:
    def _popula(self, conn):
        for mes, valor in (("2026-08", "500,00"), ("2026-09", "700,00"), ("2026-10", "900,00")):
            ledger.create_transaction(conn, {"type": "saida", "amount": valor, "occurred_on": f"{mes}-10",
                                             "description": "Mercado", "category_name": "Mercado / Alimentação",
                                             "item_name": "Arroz 5kg"})
            ledger.create_transaction(conn, {"type": "entrada", "amount": "3.000", "occurred_on": f"{mes}-05",
                                             "description": "Salário", "category_name": "Salário"})

    def test_abas_e_valores(self, conn):
        from openpyxl import load_workbook

        self._popula(conn)
        nome, conteudo = exporter.xlsx_bytes(conn, "2026-10", months=3)
        assert nome.endswith(".xlsx")
        wb = load_workbook(io.BytesIO(conteudo))
        assert wb.sheetnames == ["Lançamentos", "Resumo mensal", "Comparativo", "Itens", "Modelo"]

        lanc = wb["Lançamentos"]
        assert lanc.cell(row=1, column=1).value == "Data"
        assert lanc.max_row == 3                      # cabeçalho + 2 lançamentos de outubro
        assert lanc.cell(row=2, column=7).number_format.startswith("R$")

        resumo = wb["Resumo mensal"]
        assert resumo.cell(row=1, column=1).value == "Categoria"

        itens = wb["Itens"]
        nomes = [itens.cell(row=i, column=1).value for i in range(2, itens.max_row + 1)]
        assert "Arroz 5kg" in nomes

    def test_modelo_e_reimportavel(self, conn):
        """O modelo exportado deve ser lido pelo próprio importador (round-trip)."""
        nome, conteudo = exporter.template_bytes()
        assert nome == "modelo-lancamentos.xlsx"
        token, _ = importer.save_upload(conteudo, nome)
        previa = importer.preview(conn, token, nome, ImportOptions(default_type="saida"))
        assert previa["validas"] == len(exporter.TEMPLATE_EXAMPLES)
        assert previa["mapping"]["date"] == 0 and previa["mapping"]["amount"] == 4
        r = importer.commit(conn, token, nome, ImportOptions(default_type="saida"), previa["mapping"])
        assert r["imported"] == len(exporter.TEMPLATE_EXAMPLES)
        tipos = {t["description"]: t["type"] for t in ledger.list_transactions(conn, period="2026-10")}
        assert tipos["Salário"] == "entrada"
        assert tipos["Aluguel"] == "saida"
        assert tipos["Internet"] == "saida"

    def test_exportacao_vazia_nao_quebra(self, conn):
        from openpyxl import load_workbook

        nome, conteudo = exporter.xlsx_bytes(conn, None)
        wb = load_workbook(io.BytesIO(conteudo))
        assert wb["Lançamentos"].max_row == 1


class TestRelatorioExportavel:
    def test_matriz_usada_na_planilha(self, conn):
        ledger.create_transaction(conn, {"type": "saida", "amount": "100", "occurred_on": "2026-10-01",
                                         "description": "x", "category_name": "Lazer"})
        tabela = compare.comparison_table(conn, months=2, end_period="2026-10")
        assert tabela["periodos"] == ["2026-09", "2026-10"]
        assert tabela["linhas"][0]["name"] == "Lazer"
        assert tabela["linhas"][0]["celulas"][-1]["valor"] == 10000
