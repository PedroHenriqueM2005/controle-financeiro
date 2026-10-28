"""Testes de ponta a ponta da API HTTP e das páginas."""

from __future__ import annotations

import io

import pytest


class TestPaginas:
    @pytest.mark.parametrize(
        "caminho,trecho",
        [
            ("/", "Painel de"),
            ("/lancamentos", "Lançamentos"),
            ("/itens", "Catálogo de itens"),
            ("/relatorios", "Relatórios"),
            ("/importar", "Importar planilha"),
            ("/configuracoes", "Configurações"),
        ],
    )
    def test_paginas_renderizam(self, client, caminho, trecho):
        r = client.get(caminho)
        assert r.status_code == 200, r.text[:300]
        assert trecho in r.text
        assert "/static/app.css" in r.text

    def test_painel_com_dados(self, client):
        client.post("/api/transactions", json={
            "type": "entrada", "amount": "3.500,00", "occurred_on": "2026-10-05",
            "description": "Salário", "category_name": "Salário",
        })
        r = client.get("/?period=2026-10")
        assert r.status_code == 200
        assert "R$ 3.500,00" in r.text
        assert "<svg" in r.text          # gráficos inline
        assert "Adicionar lançamento" in r.text

    def test_periodo_invalido_cai_para_o_atual(self, client):
        assert client.get("/?period=2026-99").status_code == 200
        assert client.get("/relatorios?a=abc").status_code == 200

    def test_favicon_redireciona(self, client):
        r = client.get("/favicon.ico", follow_redirects=False)
        assert r.status_code in (302, 307)

    def test_static(self, client):
        assert client.get("/static/app.css").status_code == 200
        assert client.get("/static/app.js").status_code == 200


class TestSaude:
    def test_health(self, client):
        r = client.get("/api/health").json()
        assert r["status"] == "ok"
        assert r["app"]

    def test_healthz(self, client):
        assert client.get("/healthz").text == "ok"

    def test_openapi(self, client):
        assert client.get("/api/openapi.json").status_code == 200


class TestTransactionsAPI:
    def test_ciclo_completo(self, client):
        criado = client.post("/api/transactions", json={
            "type": "saida", "amount": "1.234,56", "occurred_on": "10/10/2026",
            "description": "Mercado", "category_name": "Mercado / Alimentação",
            "item_name": "Arroz 5kg", "method": "Pix", "recurring": True,
        })
        assert criado.status_code == 201
        tx = criado.json()
        assert tx["amount_centavos"] == 123456
        assert tx["item_name"] == "Arroz 5kg"
        tx_id = tx["id"]

        lido = client.get(f"/api/transactions/{tx_id}").json()
        assert lido["description"] == "Mercado"

        lista = client.get("/api/transactions?period=2026-10&q=arroz").json()
        assert lista["total"] == 1 and lista["items"][0]["id"] == tx_id

        editado = client.patch(f"/api/transactions/{tx_id}", json={"amount": "900"}).json()
        assert editado["amount_centavos"] == 90000
        assert editado["recurring"] == 1

        alternado = client.post(f"/api/transactions/{tx_id}/toggle").json()
        assert alternado["status"] == "previsto"

        assert client.delete(f"/api/transactions/{tx_id}").status_code == 200
        assert client.get(f"/api/transactions/{tx_id}").status_code == 404

    @pytest.mark.parametrize(
        "payload,codigo",
        [
            ({"type": "saida", "amount": "0", "occurred_on": "2026-10-01", "description": "x"}, 422),
            ({"type": "saida", "amount": "abc", "occurred_on": "2026-10-01", "description": "x"}, 422),
            ({"type": "voo", "amount": "10", "occurred_on": "2026-10-01", "description": "x"}, 422),
        ],
    )
    def test_validacao_devolve_422(self, client, payload, codigo):
        assert client.post("/api/transactions", json=payload).status_code == codigo

    def test_404(self, client):
        assert client.patch("/api/transactions/9999", json={"amount": "10"}).status_code == 404
        assert client.delete("/api/transactions/9999").status_code == 404
        assert client.post("/api/transactions/9999/toggle").status_code == 404

    def test_paginacao(self, client):
        for i in range(5):
            client.post("/api/transactions", json={
                "type": "saida", "amount": "10", "occurred_on": f"2026-10-{i + 1:02d}", "description": f"item {i}",
            })
        r = client.get("/api/transactions?period=2026-10&page=1&page_size=2").json()
        assert r["total"] == 5 and r["pages"] == 3 and len(r["items"]) == 2

    def test_periodo_invalido_no_sumario(self, client):
        assert client.get("/api/summary/2026-99").status_code == 422


class TestResumoEComparativo:
    def test_summary_com_graficos(self, client):
        client.post("/api/transactions", json={"type": "entrada", "amount": "3.000", "occurred_on": "2026-10-05",
                                               "description": "Salário", "category_name": "Salário"})
        client.post("/api/transactions", json={"type": "saida", "amount": "1.000", "occurred_on": "2026-10-06",
                                               "description": "Aluguel", "category_name": "Moradia", "item_name": "Aluguel"})
        r = client.get("/api/summary/2026-10").json()
        assert r["entradas_realizadas"] == 300000
        assert r["saidas_realizadas"] == 100000
        assert r["saldo_real"] == 200000
        assert r["salario_previsto"] == 300000          # sincronizado com o lançamento de salário
        assert r["graficos"]["serie"].startswith("<svg")
        assert r["graficos"]["donut"].startswith("<svg")
        assert isinstance(r["diario"], list) and r["diario"][0]["date"] == "2026-10-05"

    def test_compare(self, client):
        client.post("/api/transactions", json={"type": "saida", "amount": "500", "occurred_on": "2026-09-10",
                                               "description": "Mercado", "category_name": "Mercado / Alimentação"})
        client.post("/api/transactions", json={"type": "saida", "amount": "800", "occurred_on": "2026-10-10",
                                               "description": "Mercado", "category_name": "Mercado / Alimentação"})
        r = client.get("/api/compare?a=2026-10&b=2026-09").json()
        assert r["totais"]["saidas"]["delta"] == 30000
        assert r["categorias"][0]["name"] == "Mercado / Alimentação"
        assert r["graficos"]["categorias"].startswith("<svg")

    def test_compare_periodo_invalido(self, client):
        assert client.get("/api/compare?a=2026-13").status_code == 422

    def test_matrix_e_acumulado(self, client):
        client.post("/api/transactions", json={"type": "saida", "amount": "100", "occurred_on": "2026-10-10",
                                               "description": "x", "category_name": "Lazer"})
        m = client.get("/api/reports/matrix?months=6&period=2026-10").json()
        assert m["periodos"][-1] == "2026-10"
        assert any(l["name"] == "Lazer" for l in m["linhas"])
        a = client.get("/api/reports/accumulated?period=2026-10&months=6").json()
        assert a["final"] == -10000
        assert a["grafico"].startswith("<svg")


class TestCategoriasAPI:
    def test_crud(self, client):
        criada = client.post("/api/categories", json={"name": "Pets", "kind": "saida", "color": "#f472b6", "budget": 30000})
        assert criada.status_code == 201
        cat_id = criada.json()["id"]

        assert client.post("/api/categories", json={"name": "Pets"}).status_code == 422

        editada = client.patch(f"/api/categories/{cat_id}", json={"budget": 50000, "name": "Pets e vet"}).json()
        assert editada["budget_centavos"] == 50000 and editada["name"] == "Pets e vet"

        lista = client.get("/api/categories?kind=saida").json()
        assert any(c["name"] == "Pets e vet" for c in lista)

        assert client.delete(f"/api/categories/{cat_id}").json()["ok"] is True

    def test_categorias_por_mes_traz_total(self, client):
        client.post("/api/transactions", json={"type": "saida", "amount": "250", "occurred_on": "2026-10-10",
                                               "description": "feira", "category_name": "Mercado / Alimentação"})
        cats = client.get("/api/categories?kind=saida&period=2026-10").json()
        mercado = next(c for c in cats if c["name"] == "Mercado / Alimentação")
        assert mercado["total_mes"] == 25000


class TestItensAPI:
    def test_crud_e_sugestao(self, client):
        criado = client.post("/api/items", json={"name": "Café 500g", "unit": "un"})
        assert criado.status_code == 201
        item_id = criado.json()["id"]

        client.post("/api/transactions", json={"type": "saida", "amount": "21,90", "occurred_on": "2026-09-01",
                                               "description": "compra", "item_name": "Café 500g"})
        client.post("/api/transactions", json={"type": "saida", "amount": "27,90", "occurred_on": "2026-10-01",
                                               "description": "compra", "item_name": "Café 500g"})

        sugestoes = client.get("/api/items/suggest?q=cafe").json()
        assert sugestoes[0]["id"] == item_id
        assert sugestoes[0]["compras"] == 2

        item = client.get(f"/api/items/{item_id}").json()
        assert item["compras"] == 2
        assert len(item["historico"]) == 2

        check = client.get(f"/api/items/{item_id}/price-check?amount=27,90").json()
        assert check["tem_historico"] is True and check["variacao_pct"] > 0

        editado = client.patch(f"/api/items/{item_id}", json={"name": "Café torrado 500g"}).json()
        assert editado["name"] == "Café torrado 500g"

        assert client.delete(f"/api/items/{item_id}").status_code == 200
        assert client.get(f"/api/items/{item_id}").status_code == 404

    def test_item_sem_nome(self, client):
        assert client.post("/api/items", json={"name": "  "}).status_code == 422

    def test_price_check_valor_invalido(self, client):
        item = client.post("/api/items", json={"name": "X"}).json()["id"]
        assert client.get(f"/api/items/{item}/price-check?amount=abc").status_code == 422


class TestCiclosAPI:
    def test_put_e_get(self, client):
        r = client.put("/api/cycles/2026-10", json={"expected_income": 400000, "payday_day": 5, "note": "salário"})
        assert r.status_code == 200
        assert r.json()["expected_fmt"] == "R$ 4.000,00"
        assert client.get("/api/cycles/2026-10").json()["payday_day"] == 5
        assert "2026-10" in [c["period"] for c in client.get("/api/cycles").json()]

    def test_periodo_invalido(self, client):
        assert client.put("/api/cycles/2026-13", json={"expected_income": 1}).status_code == 422

    def test_repetir_recorrentes(self, client):
        client.post("/api/transactions", json={"type": "saida", "amount": "1.200", "occurred_on": "2026-09-05",
                                               "description": "Aluguel", "category_name": "Moradia", "recurring": True})
        r = client.post("/api/cycles/2026-10/repeat-recurring").json()
        assert r["criados"] == 1
        tx = client.get("/api/transactions?period=2026-10").json()["items"][0]
        assert tx["status"] == "previsto" and tx["description"] == "Aluguel"


class TestExportAPI:
    def _popula(self, client):
        client.post("/api/transactions", json={"type": "entrada", "amount": "3.000", "occurred_on": "2026-10-05",
                                               "description": "Salário", "category_name": "Salário"})
        client.post("/api/transactions", json={"type": "saida", "amount": "1.200", "occurred_on": "2026-10-06",
                                               "description": "Aluguel", "category_name": "Moradia", "item_name": "Aluguel"})

    def test_csv(self, client):
        self._popula(client)
        r = client.get("/api/export/csv?period=2026-10")
        assert r.status_code == 200
        assert "attachment" in r.headers["content-disposition"]
        assert r.content.startswith(b"\xef\xbb\xbf")            # BOM UTF-8
        texto = r.content.decode("utf-8-sig")
        assert texto.splitlines()[0].startswith("Data;Tipo;Status")
        # CSV usa ponto como separador de milhar? Não: valores "limpos" (1200,00)
        # para o Excel pt-BR interpretar como número e para reimportar sem ambiguidade.
        assert "1200,00" in texto and "3000,00" in texto

    def test_resumo_csv(self, client):
        self._popula(client)
        r = client.get("/api/export/resumo-csv?months=6")
        assert r.status_code == 200
        assert "Categoria" in r.content.decode("utf-8-sig")

    def test_xlsx_tem_abas(self, client):
        from openpyxl import load_workbook

        self._popula(client)
        r = client.get("/api/export/xlsx?period=2026-10")
        assert r.status_code == 200
        assert r.content[:2] == b"PK"
        wb = load_workbook(io.BytesIO(r.content))
        assert wb.sheetnames == ["Lançamentos", "Resumo mensal", "Comparativo", "Itens", "Modelo"]
        ws = wb["Lançamentos"]
        assert ws.cell(row=2, column=7).value in (1200.0, 3000.0)

    def test_template(self, client):
        r = client.get("/api/export/template")
        assert r.status_code == 200 and r.content[:2] == b"PK"


class TestImportAPI:
    CSV = "Data;Descrição;Categoria;Valor\n05/09/2026;Aluguel;Moradia;1.200,00\n10/09/2026;Mercado;Mercado;430,50\n"

    def test_fluxo_upload_preview_commit_undo(self, client):
        up = client.post(
            "/api/import/upload",
            files={"file": ("gastos.csv", self.CSV.encode("utf-8"), "text/csv")},
            data={"default_type": "saida"},
        )
        assert up.status_code == 200
        previa = up.json()
        assert previa["validas"] == 2 and previa["mapping"]["amount"] == 3

        corpo = {"token": previa["token"], "filename": "gastos.csv", "mapping": previa["mapping"],
                 "default_type": "saida"}
        commit = client.post("/api/import/commit", json=corpo)
        assert commit.status_code == 200
        assert commit.json()["imported"] == 2
        assert client.get("/api/transactions?period=2026-09").json()["total"] == 2

        lotes = client.get("/api/imports").json()
        assert len(lotes) == 1
        assert client.delete(f"/api/imports/{lotes[0]['id']}").json()["lancamentos_removidos"] == 2
        assert client.get("/api/transactions?period=2026-09").json()["total"] == 0

    def test_mapeamento_errado_nao_importa_nada(self, client):
        up = client.post("/api/import/upload", files={"file": ("x.csv", self.CSV.encode(), "text/csv")},
                         data={"default_type": "saida"}).json()
        corpo = {"token": up["token"], "filename": "x.csv", "mapping": {"date": 0, "amount": 99}}
        previa = client.post("/api/import/preview", json=corpo)
        assert previa.status_code == 200 and previa.json()["validas"] == 0
        # commit sem nenhuma linha válida => erro claro para o usuário
        assert client.post("/api/import/commit", json=corpo).status_code == 422

    def test_extensao_invalida(self, client):
        r = client.post("/api/import/upload", files={"file": ("x.pdf", b"%PDF-1.4", "application/pdf")})
        assert r.status_code == 422

    def test_arquivo_vazio(self, client):
        r = client.post("/api/import/upload", files={"file": ("x.csv", b"", "text/csv")})
        assert r.status_code == 422

    def test_undo_de_lote_inexistente(self, client):
        assert client.delete("/api/imports/999").status_code == 404


class TestDemo:
    def test_seed_demo(self, client):
        r = client.post("/api/demo?months=2")
        assert r.status_code == 200
        assert r.json()["lancamentos_criados"] > 20
        assert client.get("/").status_code == 200
        assert client.get("/relatorios").status_code == 200


class TestPeriodsAPI:
    def test_periods(self, client):
        client.post("/api/transactions", json={"type": "saida", "amount": "10", "occurred_on": "2026-01-05",
                                               "description": "x"})
        r = client.get("/api/periods").json()
        assert "2026-01" in r["disponiveis"]
        assert r["atual"]

    def test_methods(self, client):
        client.post("/api/transactions", json={"type": "saida", "amount": "10", "occurred_on": "2026-01-05",
                                               "description": "x", "method": "Pix"})
        assert "Pix" in client.get("/api/methods").json()


class TestPatchParcial:
    """PATCH de verdade: mandar só um campo não pode cair em 422.

    Regressão — `CategoryIn`/`ItemIn` tinham `name` obrigatório, então ajustar
    só o limite mensal (o caso mais comum em Configurações) era rejeitado.
    """

    def _categoria(self, client, nome="Pets", **extra):
        r = client.post("/api/categories", json={"name": nome, "kind": "saida", **extra})
        assert r.status_code == 201, r.text
        return r.json()["id"]

    def test_atualiza_apenas_o_limite(self, client):
        cat_id = self._categoria(client)
        r = client.patch(f"/api/categories/{cat_id}", json={"budget": 45000})
        assert r.status_code == 200, r.text
        corpo = r.json()
        assert corpo["budget_centavos"] == 45000
        assert corpo["name"] == "Pets"          # o resto permanece
        assert corpo["color"] == "#94a3b8"

    def test_limite_nulo_remove_o_limite(self, client):
        cat_id = self._categoria(client, budget=45000)
        assert client.patch(f"/api/categories/{cat_id}", json={"budget": None}).json()["budget_centavos"] is None
        orc = client.get("/api/budget/2026-10").json()
        assert "Pets" not in [c["name"] for c in orc["categorias"]]

    def test_atualiza_apenas_a_cor_e_arquiva(self, client):
        cat_id = self._categoria(client)
        assert client.patch(f"/api/categories/{cat_id}", json={"color": "#123456"}).json()["color"] == "#123456"
        arquivada = client.patch(f"/api/categories/{cat_id}", json={"archived": True}).json()
        assert arquivada["archived"] == 1
        assert all(c["id"] != cat_id for c in client.get("/api/categories?kind=saida").json())
        assert any(c["id"] == cat_id for c in client.get("/api/categories?kind=saida&include_archived=true").json())

    def test_corpo_vazio_da_422(self, client):
        cat_id = self._categoria(client)
        assert client.patch(f"/api/categories/{cat_id}", json={}).status_code == 422

    def test_categoria_inexistente_da_404(self, client):
        assert client.patch("/api/categories/999999", json={"budget": 100}).status_code == 404

    def test_nome_duplicado_da_422_e_nao_500(self, client):
        cat_id = self._categoria(client, nome="Pets")
        self._categoria(client, nome="Viagens")
        r = client.patch(f"/api/categories/{cat_id}", json={"name": "Viagens"})
        assert r.status_code == 422
        assert "já existe" in r.json()["detail"]
        # variação sem acento/maiúscula também é barrada (identidade normalizada)
        assert client.patch(f"/api/categories/{cat_id}", json={"name": "VIAGENS"}).status_code == 422

    def test_item_atualiza_somente_unidade(self, client):
        item_id = client.post("/api/items", json={"name": "Arroz 5kg", "unit": "un"}).json()["id"]
        r = client.patch(f"/api/items/{item_id}", json={"unit": "pacote"})
        assert r.status_code == 200, r.text
        assert r.json()["unit"] == "pacote"
        assert r.json()["name"] == "Arroz 5kg"

    def test_item_atualiza_somente_categoria(self, client):
        cat_id = self._categoria(client, nome="Mercearia")
        item_id = client.post("/api/items", json={"name": "Feijão 1kg"}).json()["id"]
        r = client.patch(f"/api/items/{item_id}", json={"category_id": cat_id})
        assert r.status_code == 200, r.text
        assert r.json()["category_id"] == cat_id

    def test_limite_definido_por_patch_aparece_no_orcamento(self, client):
        """Fluxo completo: ajustar o limite pela API reflete no cartão de orçamento."""
        client.post("/api/transactions", json={"type": "saida", "amount": "120,00",
                                               "occurred_on": "2026-10-04", "description": "ração",
                                               "category_name": "Pets"})
        pets = next(c for c in client.get("/api/categories?kind=saida").json() if c["name"] == "Pets")
        assert client.patch(f"/api/categories/{pets['id']}", json={"budget": 10000}).status_code == 200
        orc = client.get("/api/budget/2026-10").json()
        assert orc["orcado"] == 10000 and orc["gasto"] == 12000 and orc["acima"] is True
