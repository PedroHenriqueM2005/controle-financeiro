"""Exportação de dados em CSV e XLSX.

Detalhes que importam para planilha brasileira:

* CSV com **BOM UTF-8** e separador ``;`` — abre certo no Excel pt-BR;
* valores exportados em **reais** (não centavos) com formato de número, para
  que o Excel some/filtre corretamente;
* o XLSX sai com várias abas (lançamentos, resumo mensal, comparativo, itens) e
  uma aba **Modelo**, no mesmo formato aceito pelo importador — o ciclo
  "exportar → preencher → importar" fecha sem atrito.
"""

from __future__ import annotations

import csv
import io
import sqlite3
from datetime import datetime
from typing import Any, Sequence

from ..utils import current_period, format_date
from . import catalog, compare, ledger

CSV_HEADER = [
    "Data",
    "Tipo",
    "Status",
    "Descrição",
    "Item",
    "Categoria",
    "Valor",
    "Forma de pagamento",
    "Recorrente",
    "Origem",
    "Observação",
]

TEMPLATE_HEADER = ["Data", "Descrição", "Categoria", "Item", "Valor", "Tipo", "Status", "Forma de pagamento", "Observação"]
TEMPLATE_EXAMPLES = [
    ["05/10/2026", "Salário", "Salário", "", 3200.00, "entrada", "realizado", "Transferência", ""],
    ["05/10/2026", "Aluguel", "Moradia", "Aluguel", 1200.00, "saida", "realizado", "Boleto", "fixo"],
    ["07/10/2026", "Compra do mês", "Mercado / Alimentação", "Arroz 5kg", 28.90, "saida", "realizado", "Cartão de débito", ""],
    ["10/10/2026", "Internet", "Contas (água, luz, internet)", "Internet 300MB", 99.90, "saida", "previsto", "Débito automático", "vence dia 10"],
]


def _rows_for_export(conn: sqlite3.Connection, period: str | None, tx_type: str | None) -> list[dict[str, Any]]:
    return ledger.list_transactions(
        conn,
        period=period,
        tx_type=tx_type,
        order="asc",
        limit=None,
    )


def _row_values(tx: dict[str, Any]) -> list[Any]:
    return [
        format_date(tx["occurred_on"]),
        "Entrada" if tx["type"] == "entrada" else "Saída",
        "Previsto" if tx["status"] == "previsto" else "Realizado",
        tx["description"],
        tx.get("item_name") or "",
        tx.get("category_name") or "",
        round(tx["amount_centavos"] / 100, 2),
        tx.get("method") or "",
        "Sim" if tx.get("recurring") else "",
        tx.get("source") or "manual",
        tx.get("note") or "",
    ]


def csv_bytes(conn: sqlite3.Connection, period: str | None = None, tx_type: str | None = None) -> tuple[str, bytes]:
    """Gera o CSV de lançamentos. Retorna (nome_do_arquivo, conteúdo)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(CSV_HEADER)
    for tx in _rows_for_export(conn, period, tx_type):
        row = _row_values(tx)
        row[6] = f"{row[6]:.2f}".replace(".", ",")  # decimal brasileiro no CSV
        writer.writerow(row)

    nome = f"lancamentos-{period or 'completo'}-{datetime.now():%Y%m%d}.csv"
    conteudo = buffer.getvalue().encode("utf-8-sig")
    return nome, conteudo


def resumo_csv_bytes(conn: sqlite3.Connection, months: int = 12) -> tuple[str, bytes]:
    """CSV com a matriz categorias × meses (a "planilha mestra")."""
    tabela = compare.comparison_table(conn, months=months)
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(["Categoria"] + [c["full"].capitalize() for c in tabela["cabecalho"]] + ["Total", "Média"])
    for linha in tabela["linhas"]:
        writer.writerow(
            [linha["name"]]
            + [f'{(c["valor"] / 100):.2f}'.replace(".", ",") for c in linha["celulas"]]
            + [f'{(linha["total"] / 100):.2f}'.replace(".", ","), f'{(linha["media"] / 100):.2f}'.replace(".", ",")]
        )
    writer.writerow([])
    for chave, rotulo in (("entradas", "Entradas realizadas"), ("saidas", "Saídas realizadas"), ("saldo", "Saldo")):
        writer.writerow(
            [rotulo] + [f'{(c["valor"] / 100):.2f}'.replace(".", ",") for c in tabela["rodape"][chave]] + ["", ""]
        )
    return f"resumo-mensal-{current_period()}.csv", buffer.getvalue().encode("utf-8-sig")


def _style_header(ws, ncols: int, row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="1F2937")
    font = Font(bold=True, color="FFFFFF")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autosize(ws, larguras: Sequence[int]) -> None:
    from openpyxl.utils import get_column_letter

    for idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura


def xlsx_bytes(conn: sqlite3.Connection, period: str | None = None, months: int = 12) -> tuple[str, bytes]:
    """Pasta de trabalho completa (lançamentos + resumo + comparativo + itens + modelo)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    formato_moeda = 'R$ #,##0.00'

    # -- 1. Lançamentos -------------------------------------------------
    ws = wb.active
    ws.title = "Lançamentos"
    ws.append(CSV_HEADER)
    _style_header(ws, len(CSV_HEADER))
    txs = _rows_for_export(conn, period, None)
    for tx in txs:
        ws.append(_row_values(tx))
    for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
        for cell in row:
            cell.number_format = formato_moeda
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.font = Font(color="15803D" if cell.value == "Entrada" else "B91C1C")
    _autosize(ws, [12, 10, 11, 42, 24, 26, 14, 20, 12, 18, 30])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(CSV_HEADER))}{max(1, ws.max_row)}"

    # -- 2. Resumo mensal (categorias × meses) ---------------------------
    tabela = compare.comparison_table(conn, months=months)
    ws2 = wb.create_sheet("Resumo mensal")
    ws2.append(["Categoria"] + [c["full"].capitalize() for c in tabela["cabecalho"]] + ["Total", "Média", "Máximo"])
    _style_header(ws2, len(tabela["cabecalho"]) + 4)
    for linha in tabela["linhas"]:
        ws2.append(
            [linha["name"]]
            + [round(c["valor"] / 100, 2) for c in linha["celulas"]]
            + [round(linha["total"] / 100, 2), round(linha["media"] / 100, 2), round(linha["maximo"] / 100, 2)]
        )
    ncols = len(tabela["cabecalho"]) + 4
    for row in ws2.iter_rows(min_row=2, min_col=2, max_col=ncols):
        for cell in row:
            cell.number_format = formato_moeda
    inicio_totais = ws2.max_row + 2
    for chave, rotulo in (("entradas", "Entradas realizadas"), ("saidas", "Saídas realizadas"), ("saldo", "Saldo do mês")):
        valores = [round(c["valor"] / 100, 2) for c in tabela["rodape"][chave]]
        ws2.append([rotulo] + valores)
        for cell in ws2[ws2.max_row][1:]:
            cell.number_format = formato_moeda
        ws2.cell(row=ws2.max_row, column=1).font = Font(bold=True)
    ws2.cell(row=inicio_totais - 1, column=1).value = ""
    _autosize(ws2, [34] + [15] * len(tabela["cabecalho"]) + [15, 15, 15])
    ws2.freeze_panes = "B2"

    # -- 3. Comparativo com o mês anterior ------------------------------
    comparativo = compare.compare(conn, period or current_period())
    ws3 = wb.create_sheet("Comparativo")
    ws3.append([f"{comparativo['rotulo_atual'].capitalize()} × {comparativo['rotulo_anterior']}"])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.append([])
    ws3.append(["Categoria", "Mês atual", "Mês anterior", "Variação", "Variação %"])
    _style_header(ws3, 5, row=3)
    for linha in comparativo["categorias"]:
        ws3.append(
            [
                linha["name"],
                round(linha["atual"] / 100, 2),
                round(linha["anterior"] / 100, 2),
                round(linha["delta"] / 100, 2),
                (round(linha["pct"] / 100, 4) if linha["pct"] is not None else None),
            ]
        )
    for row in ws3.iter_rows(min_row=4, min_col=2, max_col=4):
        for cell in row:
            cell.number_format = formato_moeda
    for row in ws3.iter_rows(min_row=4, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "0,0%"
            if isinstance(cell.value, (int, float)) and cell.value > 0:
                cell.font = Font(color="B91C1C")
            elif isinstance(cell.value, (int, float)) and cell.value < 0:
                cell.font = Font(color="15803D")
    _autosize(ws3, [34, 16, 16, 16, 14])
    ws3.freeze_panes = "A4"

    # -- 4. Itens (catálogo com preço médio) -----------------------------
    ws4 = wb.create_sheet("Itens")
    ws4.append(["Item", "Categoria", "Unidade", "Compras", "Preço médio", "Último preço", "Mínimo", "Máximo", "Total gasto", "Última compra"])
    _style_header(ws4, 10)
    for item in catalog.list_items(conn, only_used=False, limit=5000):
        ws4.append(
            [
                item["name"],
                item.get("category_name") or "",
                item.get("unit") or "",
                item.get("compras") or 0,
                round((item.get("preco_medio") or 0) / 100, 2),
                round((item.get("ultimo_preco") or 0) / 100, 2),
                round((item.get("preco_min") or 0) / 100, 2),
                round((item.get("preco_max") or 0) / 100, 2),
                round((item.get("total_centavos") or 0) / 100, 2),
                format_date(item["ultima_data"]) if item.get("ultima_data") else "",
            ]
        )
    for row in ws4.iter_rows(min_row=2, min_col=5, max_col=9):
        for cell in row:
            cell.number_format = formato_moeda
    _autosize(ws4, [34, 26, 10, 10, 14, 14, 14, 14, 15, 14])
    ws4.freeze_panes = "A2"

    # -- 5. Modelo reimportável -----------------------------------------
    ws5 = wb.create_sheet("Modelo")
    ws5.append(TEMPLATE_HEADER)
    _style_header(ws5, len(TEMPLATE_HEADER))
    for exemplo in TEMPLATE_EXAMPLES:
        ws5.append(exemplo)
    for row in ws5.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = formato_moeda
    ws5.append([])
    ws5.append(["Dica: preencha a partir da linha 2, mantendo os títulos. Depois é só importar em /importar."])
    ws5.cell(row=ws5.max_row, column=1).font = Font(italic=True, color="6B7280")
    _autosize(ws5, [12, 34, 28, 22, 14, 12, 12, 22, 30])
    ws5.freeze_panes = "A2"

    stream = io.BytesIO()
    wb.save(stream)
    nome = f"financas-{period or 'completo'}-{datetime.now():%Y%m%d}.xlsx"
    return nome, stream.getvalue()


def template_bytes() -> tuple[str, bytes]:
    """Planilha modelo em branco (com exemplos) para quem prefere digitar no Excel."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    ws.append(TEMPLATE_HEADER)
    _style_header(ws, len(TEMPLATE_HEADER))
    for exemplo in TEMPLATE_EXAMPLES:
        ws.append(exemplo)
    _autosize(ws, [12, 34, 28, 22, 14, 12, 12, 22, 30])
    ws.freeze_panes = "A2"
    stream = io.BytesIO()
    wb.save(stream)
    return "modelo-lancamentos.xlsx", stream.getvalue()
