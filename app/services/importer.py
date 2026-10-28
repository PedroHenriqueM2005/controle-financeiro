"""Importação de planilhas antigas (CSV/XLSX) para dentro do banco.

Fluxo em duas etapas:

1. ``preview()``  — lê o arquivo, detecta cabeçalho e colunas, converte as
   linhas e devolve uma prévia + contagens (nada é gravado);
2. ``commit()``   — grava os lançamentos num lote (``import_batches``) que pode
   ser desfeito depois com um clique.

O reconhecimento de colunas é por *alias*: funciona tanto para a planilha
brasileira caseira (``Data;Descrição;Valor;Categoria``) quanto para extratos
bancários (``effective date / description / amount``).
"""

from __future__ import annotations

import csv
import io
import re
import sqlite3
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

from .. import config
from ..money import parse_amount
from ..utils import normalize, parse_date, period_of, truncate
from . import catalog, ledger

# ------------------------------------------------------------------ aliases

ALIASES: dict[str, tuple[str, ...]] = {
    "date": (
        "data", "date", "dia", "vencimento", "data pagamento", "data do pagamento",
        "dt", "data da transacao", "data lancamento", "effective date", "posting date",
    ),
    "amount": (
        "valor", "amount", "quantia", "value", "montante", "total", "valor r$",
        "valor (r$)", "vlr", "importancia", "importe",
    ),
    "type": (
        "tipo", "type", "natureza", "direcao", "direction", "movimento",
        "debito/credito", "d/c", "entrada/saida", "tipo lancamento",
    ),
    "description": (
        "descricao", "descrição", "description", "historico", "histórico",
        "lancamento", "lançamento", "memo", "titulo", "título", "complemento",
        "detalhe", "resumo", "narrativa",
    ),
    "category": (
        "categoria", "category", "grupo", "group", "classe", "tipo de gasto",
        "segmento", "subcategoria",
    ),
    "status": ("status", "situacao", "situação", "state", "pago?", "condicao"),
    "method": (
        "forma", "forma de pagamento", "metodo", "método", "method", "pagamento",
        "cartao", "cartão", "conta", "meio",
    ),
    "item": ("item", "produto", "product", "servico", "serviço"),
    "note": ("obs", "observacao", "observação", "note", "notas", "comentario", "comentário"),
}

FIELD_LABELS = {
    "date": "Data",
    "amount": "Valor",
    "type": "Tipo (entrada/saída)",
    "description": "Descrição",
    "category": "Categoria",
    "status": "Status (previsto/realizado)",
    "method": "Forma de pagamento",
    "item": "Item",
    "note": "Observação",
}

REQUIRED_FIELDS = ("date", "amount")

ENTRADA_KEYWORDS = {
    "entrada", "entradas", "receita", "recebimento", "credito", "credit", "c",
    "deposito", "provento", "+", "positivo", "recebido", "ganho", "salario",
}
SAIDA_KEYWORDS = {
    "saida", "despesa", "debito", "debit", "d", "pagamento", "gasto", "-",
    "negativo", "pago", "compra", "despesas",
}
PREVISTO_KEYWORDS = {
    "previsto", "prevista", "pendente", "a pagar", "a receber", "agendado",
    "planned", "pending", "scheduled", "aberto", "programado", "futuro",
}
REALIZADO_KEYWORDS = {
    "realizado", "realizada", "pago", "paga", "recebido", "recebida", "quitado",
    "quitada", "efetivado", "concluido", "done", "paid", "fechado", "liquidado",
    "compensado", "sim", "ok",
}

# pistas textuais de RECEITA, usadas quando a planilha não tem coluna "tipo".
# A categoria do lançamento tem prioridade; estas palavras são o plano B.
ENTRADA_HINTS = (
    "salario", "salario mensal", "freelance", "freela", "reembolso", "13 salario",
    "decimo terceiro", "ferias", "rescisao", "rendimento", "dividendo", "juros recebidos",
    "aluguel recebido", "recebimento", "premio", "restituicao", "valor recebido",
    "pix recebido", "transferencia recebida", "venda de",
)

MAX_ROWS = 20_000
EXCEL_EPOCH = date(1899, 12, 30)

# Para tratar um arquivo como extrato bancário (positivo = entrada) exigimos um
# volume relevante de valores negativos: poucos estornos numa planilha de gastos
# não podem inverter o sentido de todas as outras linhas.
MIN_NEGATIVOS_EXTRATO = 3
MIN_PROPORCAO_EXTRATO = 0.3


class ImportException(Exception):
    """Erro de importação amigável (arquivo inválido, coluna faltando…)."""


# ------------------------------------------------------------------ dataclasses

@dataclass
class ParsedRow:
    line: int
    occurred_on: str | None = None
    amount_centavos: int | None = None
    type: str = "saida"
    status: str = "realizado"
    description: str = ""
    category: str | None = None
    method: str | None = None
    item: str | None = None
    note: str | None = None
    raw: list[str] = field(default_factory=list)
    error: str | None = None

    @property
    def ok(self) -> bool:
        return self.error is None and bool(self.occurred_on) and bool(self.amount_centavos)

    def as_dict(self) -> dict[str, Any]:
        from ..utils import format_date

        return {
            "line": self.line,
            "date": self.occurred_on,
            "date_fmt": format_date(self.occurred_on),
            "amount": self.amount_centavos,
            "type": self.type,
            "status": self.status,
            "description": self.description,
            "category": self.category,
            "method": self.method,
            "item": self.item,
            "note": self.note,
            "error": self.error,
            "ok": self.ok,
        }

    def dedupe_key(self) -> tuple:
        return (
            self.occurred_on,
            self.amount_centavos,
            self.type,
            normalize(self.description or "")[:60],
        )


@dataclass
class ImportOptions:
    default_type: str = "saida"
    default_status: str = "realizado"
    default_category: str | None = None
    create_categories: bool = True
    create_items: bool = True
    skip_duplicates: bool = True
    sheet: int | str = 0
    header_row: int | None = None  # 0-based; None = detectar

    @classmethod
    def from_form(cls, data: dict[str, Any]) -> "ImportOptions":
        def flag(nome: str, padrao: bool = True) -> bool:
            valor = data.get(nome, padrao)
            if isinstance(valor, str):
                return valor.strip().lower() in ("1", "true", "on", "sim", "yes")
            return bool(valor) if valor is not None else padrao

        sheet = data.get("sheet", 0)
        try:
            sheet = int(sheet)
        except (TypeError, ValueError):
            sheet = 0
        header = data.get("header_row")
        try:
            header = int(header) if header not in (None, "") else None
        except (TypeError, ValueError):
            header = None

        return cls(
            default_type=data.get("default_type") if data.get("default_type") in ledger.VALID_TYPES else "saida",
            default_status=data.get("default_status") if data.get("default_status") in ledger.VALID_STATUS else "realizado",
            default_category=(data.get("default_category") or None),
            create_categories=flag("create_categories"),
            create_items=flag("create_items"),
            skip_duplicates=flag("skip_duplicates"),
            sheet=sheet,
            header_row=header,
        )


# ------------------------------------------------------------------ leitura de arquivos

def save_upload(content: bytes, filename: str) -> tuple[str, Path]:
    """Guarda o arquivo enviado numa pasta temporária e devolve (token, caminho)."""
    config.ensure_dirs()
    if len(content) > config.MAX_UPLOAD_BYTES:
        raise ImportException(
            f"Arquivo grande demais ({len(content) / 1024 / 1024:.1f} MB). "
            f"O limite é {config.MAX_UPLOAD_BYTES / 1024 / 1024:.0f} MB."
        )
    sufixo = Path(filename or "").suffix.lower()
    if sufixo not in config.ALLOWED_EXTENSIONS:
        raise ImportException(
            f"Extensão não suportada: '{sufixo}'. Envie .csv, .xlsx ou .xlsm."
        )
    token = uuid.uuid4().hex[:16]
    caminho = config.IMPORT_DIR / f"{token}{sufixo}"
    caminho.write_bytes(content)
    return token, caminho


def pending_path(token: str) -> Path:
    for sufixo in config.ALLOWED_EXTENSIONS:
        caminho = config.IMPORT_DIR / f"{token}{sufixo}"
        if caminho.exists():
            return caminho
    raise ImportException("Arquivo temporário não encontrado — envie a planilha novamente.")


def read_rows(path: str | Path, sheet: int | str = 0) -> tuple[list[list[Any]], list[str]]:
    """Lê as linhas brutas de um CSV ou XLSX. Devolve (linhas, nomes_abas)."""
    caminho = Path(path)
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(caminho, sheet)
    return _read_csv(caminho), []


def _read_csv(caminho: Path) -> list[list[Any]]:
    bruto = caminho.read_bytes()
    texto = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            texto = bruto.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if texto is None:  # pragma: no cover - latin-1 sempre decodifica
        raise ImportException("Não foi possível decodificar o arquivo CSV.")

    amostra = texto[:8192]
    try:
        dialect = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
        delimitador = dialect.delimiter
    except csv.Error:
        delimitador = ";" if amostra.count(";") >= amostra.count(",") else ","

    linhas: list[list[Any]] = []
    reader = csv.reader(io.StringIO(texto), delimiter=delimitador, skipinitialspace=True)
    for linha in reader:
        if any(str(celula).strip() for celula in linha):
            linhas.append([celula.strip() for celula in linha])
        if len(linhas) >= MAX_ROWS:
            break
    if not linhas:
        raise ImportException("O arquivo CSV está vazio.")
    return linhas


def _read_xlsx(caminho: Path, sheet: int | str = 0) -> tuple[list[list[Any]], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportException("Instale o openpyxl para importar .xlsx (pip install -r requirements.txt).") from exc

    wb = load_workbook(caminho, read_only=True, data_only=True)
    nomes = list(wb.sheetnames)
    if isinstance(sheet, str) and sheet in nomes:
        ws = wb[sheet]
    elif isinstance(sheet, int) and 0 <= sheet < len(nomes):
        ws = wb[nomes[sheet]]
    else:
        ws = wb[nomes[0]]

    linhas: list[list[Any]] = []
    for row in ws.iter_rows(values_only=True):
        if any(celula not in (None, "") for celula in row):
            linhas.append(list(row))
        if len(linhas) >= MAX_ROWS:
            break
    wb.close()
    if not linhas:
        raise ImportException("A planilha está vazia.")
    return linhas, nomes


def sheet_names(path: str | Path) -> list[str]:
    caminho = Path(path)
    if caminho.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(caminho, read_only=True)
        nomes = list(wb.sheetnames)
        wb.close()
        return nomes
    return []


# ------------------------------------------------------------------ detecção de colunas

def match_field(header: str) -> str | None:
    """Descobre qual campo lógico uma coluna representa."""
    chave = normalize(header)
    chave = re.sub(r"\s*\(.*?\)\s*", " ", chave).strip()
    if not chave:
        return None
    for field_name, aliases in ALIASES.items():
        if chave in aliases:
            return field_name
    # correspondência parcial (ex.: "valor pago" -> amount)
    for field_name, aliases in ALIASES.items():
        for alias in aliases:
            if chave == alias or chave.startswith(alias + " ") or chave.endswith(" " + alias):
                return field_name
    return None


def detect_header(rows: Sequence[Sequence[Any]], forced: int | None = None) -> tuple[int, dict[str, int]]:
    """Localiza a linha de cabeçalho e monta o mapeamento campo → coluna."""
    if forced is not None and 0 <= forced < len(rows):
        return forced, _mapping_from_row(rows[forced])

    melhor_idx, melhor_score, melhor_map = 0, -1, {}
    for idx, row in enumerate(rows[:15]):
        mapping = _mapping_from_row(row)
        score = len(mapping)
        if "date" in mapping:
            score += 2
        if "amount" in mapping:
            score += 2
        if score > melhor_score:
            melhor_idx, melhor_score, melhor_map = idx, score, mapping
    if melhor_score >= 2:
        return melhor_idx, melhor_map

    # sem cabeçalho reconhecível: deduz pelas primeiras linhas de dados
    return 0, _mapping_from_data(rows[:12])


def _mapping_from_row(row: Sequence[Any]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, celula in enumerate(row):
        field_name = match_field(str(celula or ""))
        if field_name and field_name not in mapping:
            mapping[field_name] = idx
    return mapping


def _mapping_from_data(rows: Sequence[Sequence[Any]]) -> dict[str, int]:
    """Heurística posicional para planilhas sem cabeçalho."""
    mapping: dict[str, int] = {}
    largura = max((len(r) for r in rows), default=0)
    for col in range(largura):
        valores = [r[col] for r in rows if col < len(r) and r[col] not in (None, "")]
        if not valores:
            continue
        datas = sum(1 for v in valores if _as_date(v) is not None)
        numeros = sum(1 for v in valores if parse_amount(v) is not None and _as_date(v) is None)
        if datas >= max(1, int(len(valores) * 0.6)) and "date" not in mapping:
            mapping["date"] = col
        elif numeros >= max(1, int(len(valores) * 0.6)) and "amount" not in mapping:
            mapping["amount"] = col
        elif "description" not in mapping and all(isinstance(v, str) for v in valores):
            mapping["description"] = col
    return mapping


# ------------------------------------------------------------------ conversão de valores

def _as_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20_000 < float(value) < 80_000:
        return EXCEL_EPOCH + timedelta(days=int(value))  # número de série do Excel
    return parse_date(str(value))


def _as_status(value: Any, default: str) -> str:
    chave = normalize(str(value or ""))
    if not chave:
        return default
    for alias in REALIZADO_KEYWORDS:
        if chave == alias or chave.startswith(alias):
            return "realizado"
    for alias in PREVISTO_KEYWORDS:
        if chave == alias or chave.startswith(alias):
            return "previsto"
    return default


def _as_type(value: Any) -> str | None:
    chave = normalize(str(value or "")).strip()
    if not chave:
        return None
    if chave in ENTRADA_KEYWORDS:
        return "entrada"
    if chave in SAIDA_KEYWORDS:
        return "saida"
    # prefixo só vale para alias longo: evita que "cartão" vire "crédito"
    for alias in ENTRADA_KEYWORDS:
        if len(alias) > 2 and chave.startswith(alias):
            return "entrada"
    for alias in SAIDA_KEYWORDS:
        if len(alias) > 2 and chave.startswith(alias):
            return "saida"
    return None


def entrada_category_names(conn: sqlite3.Connection) -> set[str]:
    """Nomes normalizados das categorias de entrada cadastradas."""
    return {normalize(r["name"]) for r in conn.execute("SELECT name FROM categories WHERE kind = 'entrada'")}


def _parece_entrada(descricao: str, categoria: str | None, nomes_entrada: set[str] | None) -> bool:
    """Decide se uma linha sem coluna de tipo é provavelmente uma RECEITA."""
    texto_cat = normalize(categoria)
    texto_desc = normalize(descricao)
    if nomes_entrada:
        for nome in nomes_entrada:
            if len(nome) < 4:
                continue
            if texto_cat == nome or texto_cat.startswith(nome + " ") or texto_desc == nome or texto_desc.startswith(nome + " "):
                return True
    return any(texto_desc.startswith(pista) or texto_cat.startswith(pista) for pista in ENTRADA_HINTS)


def _cell(row: Sequence[Any], mapping: dict[str, int], field_name: str) -> Any:
    idx = mapping.get(field_name)
    if idx is None or idx >= len(row):
        return None
    valor = row[idx]
    if isinstance(valor, str):
        valor = valor.strip()
    return valor if valor not in ("", None) else None


# ------------------------------------------------------------------ preview / commit

def parse_rows(
    rows: Sequence[Sequence[Any]],
    header_index: int,
    mapping: dict[str, int],
    options: ImportOptions,
    nomes_entrada: set[str] | None = None,
) -> tuple[list[ParsedRow], dict[str, Any]]:
    """Converte as linhas de dados em ``ParsedRow``, aplicando as heurísticas."""
    dados = [r for r in rows[header_index + 1:] if any(c not in (None, "") for c in r)]

    if "date" not in mapping or "amount" not in mapping:
        raise ImportException(
            "Não achei as colunas de **Data** e **Valor**. Ajuste o mapeamento manualmente na prévia."
        )

    # Heurística de extrato bancário: só entra em vigor quando NÃO há coluna de
    # tipo E os dois sinais aparecem em proporção relevante. Uma planilha de
    # gastos com um único valor negativo (estorno/ajuste) continua sendo
    # tratada pelo "tipo padrão" escolhido no formulário.
    valores = [parse_amount(_cell(r, mapping, "amount")) for r in dados]
    valores = [v for v in valores if v is not None]
    negativos = sum(1 for v in valores if v < 0)
    positivos = sum(1 for v in valores if v > 0)
    proporcao_negativos = (negativos / len(valores)) if valores else 0.0
    modo_extrato = (
        "type" not in mapping
        and positivos > 0
        and negativos >= MIN_NEGATIVOS_EXTRATO
        and MIN_PROPORCAO_EXTRATO <= proporcao_negativos <= 0.8
    )

    parsed: list[ParsedRow] = []
    for offset, row in enumerate(dados):
        numero = header_index + offset + 2  # 1-based, considerando o cabeçalho
        linha = ParsedRow(line=numero, raw=[str(c) if c is not None else "" for c in row])

        data = _as_date(_cell(row, mapping, "date"))
        if data is None:
            linha.error = "data inválida ou vazia"
            parsed.append(linha)
            continue

        bruto = _cell(row, mapping, "amount")
        centavos = parse_amount(bruto)
        if centavos is None or centavos == 0:
            linha.error = "valor inválido ou zerado"
            parsed.append(linha)
            continue

        negativo = centavos < 0
        centavos = abs(centavos)

        descricao = str(_cell(row, mapping, "description") or "").strip()
        item = str(_cell(row, mapping, "item") or "").strip() or None
        categoria_texto = _cell(row, mapping, "category")
        categoria_texto = str(categoria_texto).strip() if categoria_texto else None
        if not descricao:
            descricao = item or "Lançamento importado"

        # Ordem de decisão do tipo:
        #   1. coluna "tipo" da planilha
        #   2. valor negativo (sempre saída)
        #   3. modo extrato (positivo = entrada)
        #   4. categoria/descrição reconhecida como RECEITA
        #   5. tipo padrão escolhido no formulário
        tipo_coluna = _as_type(_cell(row, mapping, "type"))
        if tipo_coluna:
            tipo = tipo_coluna
        elif negativo:
            tipo = "saida"
        elif modo_extrato:
            tipo = "entrada"
        elif _parece_entrada(descricao, categoria_texto, nomes_entrada):
            tipo = "entrada"
        else:
            tipo = options.default_type
        if negativo:
            tipo = "saida"

        linha.occurred_on = data.isoformat()
        linha.amount_centavos = centavos
        linha.type = tipo
        linha.status = _as_status(_cell(row, mapping, "status"), options.default_status)
        linha.description = truncate(descricao, 200)
        linha.category = categoria_texto or options.default_category or None
        linha.method = str(_cell(row, mapping, "method") or "").strip() or None
        linha.item = item
        nota = str(_cell(row, mapping, "note") or "").strip() or None
        linha.note = nota

        parsed.append(linha)

    info = {
        "modo_extrato": modo_extrato,
        "proporcao_negativos": round(proporcao_negativos, 3),
        "rows_total": len(dados),
        "validas": sum(1 for p in parsed if p.ok),
        "invalidas": sum(1 for p in parsed if not p.ok),
    }
    return parsed, info


def _load_and_parse(
    path: Path, options: ImportOptions, nomes_entrada: set[str] | None = None
) -> tuple[list[list[Any]], int, dict[str, int], list[ParsedRow], dict[str, Any], list[str]]:
    rows, abas = read_rows(path, options.sheet)
    header_index, mapping = detect_header(rows, options.header_row)
    parsed, info = parse_rows(rows, header_index, mapping, options, nomes_entrada)
    return rows, header_index, mapping, parsed, info, abas


def preview(
    conn: sqlite3.Connection,
    token: str,
    filename: str,
    options: ImportOptions | None = None,
    override_mapping: dict[str, int] | None = None,
) -> dict[str, Any]:
    """Pré-visualização: mapeamento detectado, amostra de linhas e contagens."""
    opts = options or ImportOptions()
    path = pending_path(token)
    nomes = entrada_category_names(conn)
    rows, header_index, mapping, parsed, info, abas = _load_and_parse(path, opts, nomes)

    if override_mapping:
        mapping = {k: int(v) for k, v in override_mapping.items() if v not in (None, "", -1)}
        parsed, info = parse_rows(rows, header_index, mapping, opts, nomes)

    largura = max((len(r) for r in rows), default=0)
    colunas = []
    invertido = {v: k for k, v in mapping.items()}
    for idx in range(largura):
        amostras = []
        for r in rows[header_index + 1:]:
            if idx < len(r) and r[idx] not in (None, ""):
                amostras.append(str(r[idx]))
            if len(amostras) >= 3:
                break
        colunas.append(
            {
                "index": idx,
                "header": str(rows[header_index][idx]) if idx < len(rows[header_index]) else f"coluna {idx + 1}",
                "field": invertido.get(idx),
                "field_label": FIELD_LABELS.get(invertido.get(idx, ""), "— ignorar —"),
                "samples": amostras,
            }
        )

    periodos: dict[str, int] = {}
    for p in parsed:
        if p.ok:
            chave = period_of(p.occurred_on)
            periodos[chave] = periodos.get(chave, 0) + 1

    duplicadas = _count_duplicates(conn, parsed, opts)

    categorias_novas = sorted(
        {
            p.category
            for p in parsed
            if p.ok and p.category and not ledger.ensure_category(conn, p.category, p.type, create=False)
        }
    )
    itens_novos = sorted(
        {
            (p.item or p.description)
            for p in parsed
            if p.ok and p.type == "saida" and opts.create_items
            and not catalog.ensure_item(conn, p.item or p.description, create=False)
        }
    )

    return {
        "token": token,
        "filename": filename,
        "sheets": abas,
        "sheet": opts.sheet,
        "header_index": header_index,
        "columns": colunas,
        "mapping": mapping,
        "rows_total": info["rows_total"],
        "validas": info["validas"],
        "invalidas": info["invalidas"],
        "duplicadas": duplicadas,
        "modo_extrato": info["modo_extrato"],
        "proporcao_negativos": info.get("proporcao_negativos"),
        "periodos": dict(sorted(periodos.items(), reverse=True)),
        "periodos_count": len(periodos),
        "primeiro_periodo": min(periodos) if periodos else None,
        "ultimo_periodo": max(periodos) if periodos else None,
        "categorias_novas": categorias_novas[:30],
        "itens_novos": itens_novos[:30],
        "preview": [p.as_dict() for p in parsed[:15]],
        "errors": [
            {"row": p.line, "error": p.error, "raw": truncate(" | ".join(p.raw), 120)}
            for p in parsed
            if not p.ok
        ][:15],
        "options": {
            "default_type": opts.default_type,
            "default_status": opts.default_status,
            "default_category": opts.default_category,
            "create_categories": opts.create_categories,
            "create_items": opts.create_items,
            "skip_duplicates": opts.skip_duplicates,
        },
    }


def _existing_keys(conn: sqlite3.Connection, periods: Iterable[str]) -> set[tuple]:
    periodos = sorted(set(periods))
    if not periodos:
        return set()
    marcadores = ",".join("?" * len(periodos))
    chaves = set()
    for r in conn.execute(
        f"""
        SELECT t.occurred_on, t.amount_centavos, t.type, t.description
        FROM transactions t WHERE t.period IN ({marcadores})
        """,
        tuple(periodos),
    ):
        chaves.add((r["occurred_on"], int(r["amount_centavos"]), r["type"], normalize(r["description"] or "")[:60]))
    return chaves


def _count_duplicates(conn: sqlite3.Connection, parsed: Sequence[ParsedRow], options: ImportOptions) -> int:
    if not options.skip_duplicates:
        return 0
    existentes = _existing_keys(conn, [period_of(p.occurred_on) for p in parsed if p.ok])
    vistos: set[tuple] = set()
    duplicadas = 0
    for p in parsed:
        if not p.ok:
            continue
        chave = p.dedupe_key()
        if chave in existentes or chave in vistos:
            duplicadas += 1
        vistos.add(chave)
    return duplicadas


def commit(
    conn: sqlite3.Connection,
    token: str,
    filename: str,
    options: ImportOptions | None = None,
    override_mapping: dict[str, int] | None = None,
) -> dict[str, Any]:
    """Grava a importação como um lote rastreável (e reversível)."""
    opts = options or ImportOptions()
    path = pending_path(token)
    nomes = entrada_category_names(conn)
    rows, header_index, mapping, parsed, info, _ = _load_and_parse(path, opts, nomes)
    if override_mapping:
        mapping = {k: int(v) for k, v in override_mapping.items() if v not in (None, "", -1)}
        parsed, info = parse_rows(rows, header_index, mapping, opts, nomes)

    source_tag = f"import:{filename}:{token[:8]}"
    cur = conn.execute(
        """
        INSERT INTO import_batches (filename, source_tag, rows_total, rows_imported, rows_skipped)
        VALUES (?, ?, ?, 0, 0)
        """,
        (filename, source_tag, info["rows_total"]),
    )
    batch_id = int(cur.lastrowid)

    existentes = _existing_keys(conn, [period_of(p.occurred_on) for p in parsed if p.ok])
    vistos: set[tuple] = set()
    importados = pulados = 0
    datas: list[str] = []

    for p in parsed:
        if not p.ok:
            pulados += 1
            continue
        chave = p.dedupe_key()
        if opts.skip_duplicates and (chave in existentes or chave in vistos):
            pulados += 1
            continue
        vistos.add(chave)

        categoria_id = None
        if p.category:
            categoria_id = ledger.ensure_category(
                conn, p.category, p.type, create=opts.create_categories
            )
            if categoria_id is None:
                categoria_id = ledger.ensure_category(conn, opts.default_category or "Outras saídas", p.type)

        item_id = None
        if p.type == "saida":
            nome_item = p.item or (p.description if len(p.description) <= 60 else None)
            if nome_item:
                item_id = catalog.ensure_item(conn, nome_item, categoria_id, create=opts.create_items)

        conn.execute(
            """
            INSERT INTO transactions
                (type, status, occurred_on, period, amount_centavos, description,
                 category_id, item_id, method, recurring, source, import_id, note)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?)
            """,
            (
                p.type, p.status, p.occurred_on, period_of(p.occurred_on), p.amount_centavos,
                p.description, categoria_id, item_id, p.method, source_tag, batch_id, p.note,
            ),
        )
        importados += 1
        datas.append(p.occurred_on)

    conn.execute(
        """
        UPDATE import_batches SET rows_imported = ?, rows_skipped = ?, first_date = ?, last_date = ?
        WHERE id = ?
        """,
        (importados, pulados, min(datas) if datas else None, max(datas) if datas else None, batch_id),
    )

    for periodo in {period_of(d) for d in datas}:
        ledger.sync_cycle_from_salary(conn, periodo)

    try:
        path.unlink()  # o arquivo temporário já cumpriu seu papel
    except OSError:  # pragma: no cover
        pass

    return {
        "batch_id": batch_id,
        "filename": filename,
        "source_tag": source_tag,
        "rows_total": info["rows_total"],
        "imported": importados,
        "skipped": pulados,
        "invalidas": info["invalidas"],
        "periodos": sorted({period_of(d) for d in datas}),
        "primeiro": min(datas) if datas else None,
        "ultimo": max(datas) if datas else None,
    }


def discard_pending(token: str) -> None:
    try:
        pending_path(token).unlink()
    except (ImportException, OSError):
        pass
